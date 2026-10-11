import re
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook

from common import csv_file, main

ROW_ID = re.compile(r"[^\s`|]+\.xlsx\|Data to publish\|\d+")


def amount(value):
    try:
        return Decimal(str(value).replace(",", ""))
    except InvalidOperation as exc:
        raise ValueError(f"invalid amount: {value}") from exc


def check(root):
    out = root / "outputs"
    clean = csv_file(out / "analytics-full" / "spend-clean.csv")
    indexed = {row["row_id"]: row for row in clean}
    issues = []
    if len(indexed) != len(clean):
        issues.append("analytics-full/spend-clean.csv has duplicate row IDs")
    sample_files = sorted((out / "samples").glob("*.csv"))
    if not sample_files:
        return issues + ["no sample CSV files"]
    sampled = []
    for path in sample_files:
        rows = csv_file(path)
        if not rows:
            issues.append(f"{path.name}: empty sample")
        for row in rows:
            sampled.append((path.name, row))
    program = (out / "audit-program.md").read_text(encoding="utf-8")
    for row_id in ROW_ID.findall(program):
        sampled.append(("audit-program.md", {"row_id": row_id}))
    workbooks = {}
    try:
        for origin, row in sampled:
            row_id = row.get("row_id", "")
            expected = indexed.get(row_id)
            if expected is None:
                issues.append(f"{origin}: row ID missing from analytics-full/spend-clean.csv: {row_id}")
                continue
            for field in ("supplier_normalized", "Supplier name", "net_amount", "Net amount"):
                if field in row and row[field] != "":
                    if field.lower().endswith("amount") or field == "net_amount":
                        if amount(row[field]) != amount(expected[field]):
                            issues.append(f"{origin}: clean amount mismatch: {row_id}")
                    elif row[field] != expected[field]:
                        issues.append(f"{origin}: clean supplier mismatch: {row_id}")
            parts = row_id.rsplit("|", 2)
            if len(parts) != 3 or parts[1] != "Data to publish" or not parts[2].isdigit():
                issues.append(f"{origin}: invalid source row ID: {row_id}")
                continue
            filename, sheet, sheet_row = parts[0], parts[1], int(parts[2])
            if filename != expected.get("source_file") or str(sheet_row) != expected.get("source_sheet_row"):
                issues.append(f"{origin}: source pointer differs from clean spend: {row_id}")
                continue
            if filename not in workbooks:
                workbooks[filename] = load_workbook(root / "data/spend" / filename, read_only=True, data_only=True)
            worksheet = workbooks[filename][sheet]
            columns = None
            for cells in worksheet.iter_rows(min_row=1, max_row=min(10, worksheet.max_row), values_only=True):
                if {"Supplier name", "Net amount"}.issubset(set(cells)):
                    columns = {str(value): index for index, value in enumerate(cells) if value is not None}
                    break
            if columns is None:
                issues.append(f"{filename}: source header missing")
                continue
            values = next(worksheet.iter_rows(min_row=sheet_row, max_row=sheet_row, values_only=True))
            supplier = values[columns["Supplier name"]]
            source_amount = values[columns["Net amount"]]
            if str(supplier).strip() != expected["Supplier name"].strip() or amount(source_amount) != amount(expected["net_amount"]):
                issues.append(f"{origin}: workbook supplier or amount mismatch: {row_id}")
    finally:
        for workbook in workbooks.values():
            workbook.close()
    return issues


if __name__ == "__main__":
    raise SystemExit(main(check))
